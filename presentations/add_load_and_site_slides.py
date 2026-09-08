"""Insert the two new result slides into presentations/BigData_ITMO.pptx.

    python presentations/build_bigdata_deck_assets.py --only load
    python presentations/build_bigdata_deck_assets.py --only site
    python presentations/add_load_and_site_slides.py

Slide 11  "Part 5 - It was measuring start-up"        -> slide11_load_experiment.png
Slide 12  "Part 6 - The overall effect on the site"   -> slide12_site_effect.png

Both are built by cloning slide 10 ("Part 4 - PySpark: Windows vs Cluster"): the
same ``5_Custom Layout``, the same white rounded card behind the figure, the same
title run properties (Arial bold 26 pt, #121212, left aligned).  Cloning rather
than re-creating is deliberate - the title styling lives in ``a:pPr/a:defRPr`` on
the slide, not in the layout, so a from-scratch placeholder would not match.

The script is idempotent: if a slide with either title already exists it is left
alone, so re-running after regenerating the PNGs is safe.  Close PowerPoint
first - it holds the pre-edit copy in memory and will clobber the write.
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "presentations" / "BigData_ITMO.pptx"
ASSETS = ROOT / "presentations" / "assets" / "bigdata_itmo"

# The slide we clone, and where the two new ones go (1-based, after insertion).
TEMPLATE_TITLE = "Part 4 — PySpark: Windows vs Cluster"

NEW_SLIDES = [
    {
        "title": "Part 5 — It was measuring start-up",
        "image": ASSETS / "slide11_load_experiment.png",
        "position": 11,
    },
    {
        "title": "Part 6 — The overall effect on the site",
        "image": ASSETS / "slide12_site_effect.png",
        "position": 12,
    },
]


def _find_slide(prs, title: str):
    for slide in prs.slides:
        if slide.shapes.title is not None and slide.shapes.title.text_frame.text.strip() == title:
            return slide
    return None


def _title_sp(template_slide):
    for shape in template_slide.shapes:
        if shape == template_slide.shapes.title:
            return shape._element
    raise LookupError("template slide has no title placeholder")


def _card_sp(template_slide):
    """The white rounded rectangle that sits behind the figure."""
    for shape in template_slide.shapes:
        if shape.shape_type is not None and shape.has_text_frame and not shape.is_placeholder:
            return shape._element
    raise LookupError("template slide has no rounded card")


def _set_title_text(sp, text: str) -> None:
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    runs = sp.findall(f".//{ns}r")
    if not runs:
        raise LookupError("cloned title has no run to retext")
    runs[0].find(f"{ns}t").text = text
    for extra in runs[1:]:
        extra.getparent().remove(extra)


def _move_slide(prs, slide, position: int) -> None:
    """Move ``slide`` to 1-based ``position`` by reordering p:sldIdLst."""
    id_list = prs.slides._sldIdLst
    entries = list(id_list)
    rid = slide.part.partname  # only used for the error message
    for entry in entries:
        if prs.part.related_part(entry.rId) is slide.part:
            id_list.remove(entry)
            id_list.insert(position - 1, entry)
            return
    raise LookupError(f"slide {rid} not found in sldIdLst")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--deck", default=str(DECK),
                        help="Deck to edit (default: the project deck).")
    args = parser.parse_args(argv)
    deck = Path(args.deck)

    if not deck.exists():
        print(f"deck not found: {deck}")
        return 1

    lock = deck.with_name("~$" + deck.name)
    if lock.exists():
        print(f"REFUSING: {lock.name} exists - the deck is open in PowerPoint. Close it first.")
        return 1

    missing = [s["image"] for s in NEW_SLIDES if not s["image"].exists()]
    if missing:
        for path in missing:
            print(f"missing figure: {path}")
        print("run: python presentations/build_bigdata_deck_assets.py --only load  (and --only site)")
        return 1

    prs = Presentation(str(deck))
    template = _find_slide(prs, TEMPLATE_TITLE)
    if template is None:
        print(f"template slide not found: {TEMPLATE_TITLE!r}")
        return 1

    title_proto = _title_sp(template)
    card_proto = _card_sp(template)
    picture = next(sh for sh in template.shapes
                   if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    pic_pos = (picture.left, picture.top, picture.width, picture.height)

    added = 0
    for spec in NEW_SLIDES:
        if _find_slide(prs, spec["title"]) is not None:
            print(f"already present, skipping: {spec['title']}")
            continue

        slide = prs.slides.add_slide(template.slide_layout)
        # Drop every placeholder the layout supplied; the clone brings its own.
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)

        title_sp = copy.deepcopy(title_proto)
        _set_title_text(title_sp, spec["title"])
        slide.shapes._spTree.append(title_sp)

        # The card must precede the picture so it renders behind it.
        slide.shapes._spTree.append(copy.deepcopy(card_proto))

        slide.shapes.add_picture(str(spec["image"]), *pic_pos)

        _move_slide(prs, slide, spec["position"])
        added += 1
        print(f"added slide {spec['position']}: {spec['title']}")

    if not added:
        print("nothing to do")
        return 0

    prs.save(str(deck))
    print(f"saved {deck} ({len(prs.slides)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
